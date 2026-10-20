import json

import pytest

from attic.loaders import LoadError, load_text


def test_utf8_and_windows_turkish_text(tmp_path):
    a = tmp_path / "a.txt"
    a.write_text("çalışma notları", encoding="utf-8")
    b = tmp_path / "b.txt"
    b.write_bytes("ödev ğüşı".encode("cp1254"))
    assert load_text(a) == "çalışma notları"
    assert load_text(b) == "ödev ğüşı"


def test_binary_disguised_as_text_fails(tmp_path):
    p = tmp_path / "x.txt"
    p.write_bytes(b"\x00\x01\x02binary")
    with pytest.raises(LoadError):
        load_text(p)


def test_html_drops_scripts(tmp_path):
    p = tmp_path / "page.html"
    p.write_text("<html><script>var x=1</script><p>Hello</p><p>World</p></html>")
    text = load_text(p)
    assert "Hello" in text and "World" in text and "var x" not in text


def test_notebook_cells(tmp_path):
    p = tmp_path / "nb.ipynb"
    p.write_text(json.dumps({"cells": [
        {"cell_type": "markdown", "source": ["# Churn model\n"]},
        {"cell_type": "code", "source": "import pandas as pd"},
    ]}))
    assert load_text(p) == "# Churn model\n\nimport pandas as pd"


def test_office_formats(tmp_path):
    import docx
    import openpyxl
    from pptx import Presentation

    d = docx.Document()
    d.add_paragraph("Internship report draft")
    d.save(tmp_path / "r.docx")
    assert "Internship report draft" in load_text(tmp_path / "r.docx")

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quarterly plan"
    prs.save(tmp_path / "s.pptx")
    assert "Slide 1" in load_text(tmp_path / "s.pptx")
    assert "Quarterly plan" in load_text(tmp_path / "s.pptx")

    wb = openpyxl.Workbook()
    wb.active.append(["item", "cost"])
    wb.active.append(["server", 12])
    wb.save(tmp_path / "b.xlsx")
    assert "server\t12" in load_text(tmp_path / "b.xlsx")


def test_corrupt_pdf_is_a_load_error_not_a_crash(tmp_path):
    p = tmp_path / "broken.pdf"
    p.write_bytes(b"%PDF-1.4 not really")
    with pytest.raises(LoadError):
        load_text(p)
